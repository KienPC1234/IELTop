import sys
import re
from pathlib import Path
import docx
import pptx

sys.stdout.reconfigure(encoding='utf-8')
root = Path(r'C:\Users\kienp\Downloads\Lesson Content-20260915T152353Z-1-001\Lesson Content')

url_pattern = re.compile(r'https?://[^\s"\'<>]+', re.IGNORECASE)
cam_pattern = re.compile(r'(cambridge|cam\s*\d|ielts|study4|test\s*\d)', re.IGNORECASE)

found_urls = []
found_refs = []

for p in sorted(root.rglob('*')):
    rel = p.relative_to(root).as_posix()
    if p.suffix.lower() == '.docx':
        try:
            doc = docx.Document(str(p))
            # Check hyperlinks in docx part
            for r in doc.part.rels.values():
                if 'hyperlink' in r.reltype:
                    found_urls.append((rel, [r.target_ref]))

            for para in doc.paragraphs:
                t = para.text.strip()
                if not t:
                    continue
                urls = url_pattern.findall(t)
                if urls:
                    found_urls.append((rel, urls))
                if cam_pattern.search(t):
                    found_refs.append((rel, t[:160]))
        except Exception as ex:
            pass
    elif p.suffix.lower() == '.pptx':
        try:
            prs = pptx.Presentation(str(p))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            urls = url_pattern.findall(t)
                            if urls:
                                found_urls.append((rel, urls))
                            if 'study4' in t.lower():
                                found_refs.append((rel, t[:160]))
        except Exception:
            pass

print(f"=== Found URLs ({len(found_urls)}) ===")
seen = set()
for f, urls in found_urls:
    for u in urls:
        if u not in seen:
            seen.add(u)
            print(f"{f}: {u}")

print(f"\n=== References ({len(found_refs)}) ===")
for f, t in found_refs[:40]:
    print(f"{f}: {t}")
